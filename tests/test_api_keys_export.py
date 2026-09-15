"""Testy Zadania 1 (samoobsługowe klucze API, `POST /api/v1/keys/request`) i
Zadania 3 (eksport na żądanie, `GET /api/v1/export/{format}`) — offline, bez
połączenia do prawdziwej bazy, ten sam wzorzec atrap co `tests/test_api.py`
(zob. jego nagłówek). `auth_pool`/`data_pool` podmienione na lekkie atrapy w
pamięci; `api.rate_limit.limiter` zresetowany przed każdym testem, żeby stan
slowapi (in-memory, per-proces) nie przeciekał między testami korzystającymi
z tego samego adresu IP klienta testowego."""

from contextlib import asynccontextmanager
from datetime import UTC, datetime

import openpyxl
import pytest
from fastapi.testclient import TestClient

from api import db, main
from api.rate_limit import limiter
from api.routers import keys as keys_router


class FakeAuthPool:
    """`auth_pool()` atrapa — `keys.py`/`export.py` wołają `fetchval`/`execute`
    bezpośrednio na puli (bez `.acquire()`), więc atrapa tylko tego potrzebuje."""

    def __init__(self):
        self.api_keys: list[dict] = []
        self.export_requests: list[dict] = []
        self.db_access_requests: list[dict] = []

    async def fetchval(self, query, *args):
        if "FROM api_keys" in query:
            ip, window_start = args
            return sum(
                1 for r in self.api_keys if r["created_from_ip"] == ip and r["created_at"] > window_start
            )
        if "FROM export_requests" in query:
            ip, window_start = args
            return sum(
                1 for r in self.export_requests if r["ip"] == ip and r["created_at"] > window_start
            )
        if "FROM db_access_requests" in query:
            ip, window_start = args
            return sum(
                1
                for r in self.db_access_requests
                if r["requested_from_ip"] == ip and r["requested_at"] > window_start
            )
        raise AssertionError(f"Nieoczekiwane zapytanie w atrapie: {query}")

    async def execute(self, query, *args):
        if "INSERT INTO api_keys" in query:
            key_hash, owner_label, requester_label, created_from_ip, rate_limit = args
            self.api_keys.append(
                {
                    "key_hash": key_hash,
                    "owner_label": owner_label,
                    "requester_label": requester_label,
                    "self_service": True,
                    "created_from_ip": created_from_ip,
                    "rate_limit_per_hour": rate_limit,
                    "created_at": datetime.now(UTC),
                }
            )
        elif "INSERT INTO export_requests" in query:
            (ip,) = args
            self.export_requests.append({"ip": ip, "created_at": datetime.now(UTC)})
        elif "INSERT INTO db_access_requests" in query:
            requester_label, requested_from_ip = args
            self.db_access_requests.append(
                {
                    "requester_label": requester_label,
                    "requested_from_ip": requested_from_ip,
                    "requested_at": datetime.now(UTC),
                }
            )
        else:
            raise AssertionError(f"Nieoczekiwane zapytanie w atrapie: {query}")


class FakeDataConn:
    def __init__(self):
        self.calls: list[tuple[str, tuple]] = []

    async def fetch(self, sql, *params):
        self.calls.append((sql, params))
        return []


class FakeDataPool:
    def __init__(self):
        self.conn = FakeDataConn()

    def acquire(self):
        pool = self

        @asynccontextmanager
        async def _acquire():
            yield pool.conn

        return _acquire()


@pytest.fixture
def client(monkeypatch):
    async def _noop_init():
        return None

    async def _noop_close():
        return None

    monkeypatch.setattr(main, "init_pools", _noop_init)
    monkeypatch.setattr(main, "close_pools", _noop_close)
    limiter.reset()
    with TestClient(main.app) as c:
        yield c
    limiter.reset()


@pytest.fixture
def fake_auth_pool(monkeypatch):
    pool = FakeAuthPool()
    monkeypatch.setattr(db, "_auth_pool", pool)
    return pool


@pytest.fixture
def fake_data_pool(monkeypatch):
    pool = FakeDataPool()
    monkeypatch.setattr(db, "_data_pool", pool)
    return pool


# --- Zadanie 1 — samoobsługowe klucze ---------------------------------


def test_get_client_ip_prefers_fly_client_ip_header(client, fake_auth_pool):
    """Fly.io stawia realny adres klienta w `Fly-Client-IP` -
    `request.client.host` na Fly to zawsze wewnętrzny adres proxy (odkryte
    tej sesji: 3 klucze testowe z tego samego 172.16.x.x), więc limit per-IP
    musiałby dzielić pulę między wszystkich odwiedzających bez tej poprawki."""
    resp = client.post(
        "/api/v1/keys/request",
        json={"label": "fly-header-test"},
        headers={"Fly-Client-IP": "203.0.113.5"},
    )
    assert resp.status_code == 200
    assert fake_auth_pool.api_keys[0]["created_from_ip"] == "203.0.113.5"


def test_get_client_ip_falls_back_to_x_forwarded_for(client, fake_auth_pool):
    resp = client.post(
        "/api/v1/keys/request",
        json={"label": "xff-test"},
        headers={"X-Forwarded-For": "198.51.100.7, 10.0.0.1"},
    )
    assert resp.status_code == 200
    assert fake_auth_pool.api_keys[0]["created_from_ip"] == "198.51.100.7"


def test_ip_signup_limit_enforced(client, fake_auth_pool):
    """4. próba klucza z tego samego IP w ciągu doby zwraca 429."""
    for i in range(3):
        resp = client.post("/api/v1/keys/request", json={"label": f"student-{i}"})
        assert resp.status_code == 200, resp.text

    resp = client.post("/api/v1/keys/request", json={"label": "student-4"})
    assert resp.status_code == 429


def test_self_service_key_has_correct_rate_limit(client, fake_auth_pool):
    """Nowo wygenerowany klucz self-service ma rate_limit_per_hour=60."""
    resp = client.post("/api/v1/keys/request", json={"label": "jane-doe"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["rate_limit_per_hour"] == 60
    assert body["api_key"].startswith("nfg_edu_")
    assert fake_auth_pool.api_keys[0]["rate_limit_per_hour"] == 60
    assert fake_auth_pool.api_keys[0]["self_service"] is True


def test_signup_endpoint_itself_rate_limited(client, fake_auth_pool, monkeypatch):
    """Szybkie 10 żądań pod rząd na /keys/request - część odrzucona.

    Podnosi limit dobowy per-IP (Zadanie 1c) wysoko, żeby izolować warstwę
    testowaną tutaj (ogólny limit 5/minutę na sam endpoint, Zadanie 1c
    'niezależny od limitu dobowego') od tej sprawdzanej osobno w
    `test_ip_signup_limit_enforced`."""
    monkeypatch.setattr(keys_router, "IP_SIGNUP_MAX_PER_DAY", 1000)

    statuses = [
        client.post("/api/v1/keys/request", json={"label": f"flood-{i}"}).status_code
        for i in range(10)
    ]
    assert 429 in statuses
    assert statuses.count(200) <= 5


def test_signup_requires_label(client, fake_auth_pool):
    resp = client.post("/api/v1/keys/request", json={"label": "a"})
    assert resp.status_code == 422


# --- Zadanie 3 — eksport na żądanie ------------------------------------


def test_export_csv_returns_file_without_key(client, fake_auth_pool, fake_data_pool):
    """/export/csv działa bez klucza API."""
    resp = client.get("/api/v1/export/csv")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/zip"


def test_export_excel_returns_file_without_key(client, fake_auth_pool, fake_data_pool):
    """/export/excel działa bez klucza API (osobny test od CSV - jedno IP
    ma prawo do 1 eksportu/5 min, zob. `test_export_rate_limit_enforced`)."""
    resp = client.get("/api/v1/export/excel")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]


def test_export_default_range_is_full_history(client, fake_auth_pool, fake_data_pool):
    """Brak from/to = dane od 2019-01 do dziś."""
    resp = client.get("/api/v1/export/csv")
    assert resp.status_code == 200
    disposition = resp.headers["content-disposition"]
    today_month = datetime.now(UTC).strftime("%Y-%m")
    assert "2019-01" in disposition
    assert today_month in disposition

    from_param, to_param = fake_data_pool.conn.calls[0][1][:2]
    assert from_param.isoformat() == "2019-01-01"
    assert to_param <= datetime.now(UTC).date()


def test_export_custom_range_filters_correctly(client, fake_auth_pool, fake_data_pool):
    """from=2023-01&to=2025-12 zwraca tylko dane z tego okresu."""
    resp = client.get("/api/v1/export/csv?from=2023-01&to=2025-12")
    assert resp.status_code == 200
    assert "norfingen_2023-01_2025-12.zip" in resp.headers["content-disposition"]

    from_param, to_param = fake_data_pool.conn.calls[0][1][:2]
    assert from_param.isoformat() == "2023-01-01"
    assert to_param.isoformat() == "2025-12-31"


def test_export_rate_limit_enforced(client, fake_auth_pool, fake_data_pool):
    """2. żądanie eksportu w ciągu 5 minut z tego samego IP zwraca 429."""
    first = client.get("/api/v1/export/csv")
    assert first.status_code == 200
    second = client.get("/api/v1/export/csv")
    assert second.status_code == 429


def test_export_invalid_date_range_returns_400(client, fake_auth_pool, fake_data_pool):
    """from > to zwraca czytelny błąd, nie 500."""
    resp = client.get("/api/v1/export/csv?from=2025-12&to=2023-01")
    assert resp.status_code == 400
    # brak wpisu w liczniku eksportów - żądanie odrzucone przed enforce_export_rate_limit
    assert fake_auth_pool.export_requests == []


def test_export_rejects_unknown_format(client, fake_auth_pool, fake_data_pool):
    resp = client.get("/api/v1/export/pdf")
    assert resp.status_code == 404


def test_excel_export_produces_valid_workbook_with_expected_sheets():
    """Verifica offline pura funkcji generującej bajty (bez FastAPI/HTTP) -
    upewnia się, że nagłówki arkuszy są obecne nawet przy pustym wyniku
    zapytań (fallback `_COLUMNS`, zob. `api/export_data.py`)."""
    import asyncio
    from datetime import date

    from api import export_data

    class _EmptyConn:
        async def fetch(self, sql, *params):
            return []

    class _EmptyPool:
        def acquire(self):
            @asynccontextmanager
            async def _acquire():
                yield _EmptyConn()

            return _acquire()

    import api.db as db_module

    original = db_module._data_pool
    db_module._data_pool = _EmptyPool()
    try:
        raw = asyncio.run(export_data.generate_excel_export_bytes(date(2019, 1, 1), date.today()))
    finally:
        db_module._data_pool = original

    import io

    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert set(wb.sheetnames) == {
        "PL_miesiecznie",
        "Faktury_sprzedazy",
        "Faktury_zakupu",
        "Payroll_miesiecznie",
        "Payroll_per_pracownik",
        "Postingi_GL",
    }
    ws = wb["PL_miesiecznie"]
    assert [c.value for c in next(ws.iter_rows(max_row=1))] == [
        "month",
        "revenue",
        "labor_cost",
        "operating_cost",
        "cogs",
        "operating_result",
    ]


# --- Portal, ścieżka 3 — żywy dostęp do bazy na żądanie ----------------
#
# `test_portal_reader_cannot_write` i `test_portal_reader_statement_timeout_enforced`
# (nazwy z promptu) wymagają z definicji prawdziwego połączenia do bazy jako
# `portal_reader` - nie dają się sensownie zamockować bez sprawdzania czegoś
# innego niż to, co mają sprawdzić (uprawnienia GRANT/RLS i realny
# `statement_timeout` egzekwowane przez Postgres, nie przez appkę). Zgodnie z
# konwencją reszty pakietu (offline, zob. nagłówek tego pliku) - te dwa testy
# żyją jako osobny, ręcznie uruchamiany skrypt (`scripts/verify_portal_reader_readonly.py`),
# NIE w pytest. Wynik uruchomienia (dokładne komunikaty błędów) jest
# udokumentowany w docs/SESSION_HANDOFF.md, zgodnie z wymogiem promptu
# "RĘCZNIE zweryfikowany test... nie tylko test jednostkowy".


@pytest.fixture
def portal_env(monkeypatch):
    monkeypatch.setenv("PORTAL_DB_HOST", "db.example.supabase.co")
    monkeypatch.setenv("PORTAL_READER_USERNAME", "portal_reader.testref")
    monkeypatch.setenv("PORTAL_READER_PASSWORD", "test-password-123")


def test_db_access_request_returns_credentials(client, fake_auth_pool, portal_env):
    """POST /db-access/request zwraca kompletne dane połączenia."""
    resp = client.post("/api/v1/db-access/request", json={"label": "power-bi-test"})
    assert resp.status_code == 200
    body = resp.json()
    assert body["host"] == "db.example.supabase.co"
    assert body["port"] == 5432
    assert body["database"] == "postgres"
    assert body["username"] == "portal_reader.testref"
    assert body["password"] == "test-password-123"
    assert "read-only" in body["message"].lower()
    assert fake_auth_pool.db_access_requests[0]["requester_label"] == "power-bi-test"


def test_db_access_missing_env_returns_503(client, fake_auth_pool, monkeypatch):
    """Brak skonfigurowanych sekretów -> 503 czytelny, nie 500/KeyError."""
    monkeypatch.delenv("PORTAL_DB_HOST", raising=False)
    monkeypatch.delenv("PORTAL_READER_USERNAME", raising=False)
    monkeypatch.delenv("PORTAL_READER_PASSWORD", raising=False)
    resp = client.post("/api/v1/db-access/request", json={"label": "no-env-test"})
    assert resp.status_code == 503


def test_db_access_ip_limit_enforced(client, fake_auth_pool, portal_env):
    """4. żądanie z tego samego IP w ciągu doby zwraca 429."""
    for i in range(3):
        resp = client.post("/api/v1/db-access/request", json={"label": f"bi-user-{i}"})
        assert resp.status_code == 200, resp.text

    resp = client.post("/api/v1/db-access/request", json={"label": "bi-user-4"})
    assert resp.status_code == 429
